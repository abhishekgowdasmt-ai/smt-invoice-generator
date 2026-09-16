"""Parse all RAC duty Excel workbooks into history JSON + unique drivers."""
import json
import os
import re
from collections import Counter, defaultdict
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

SRC_DIR = r'C:\Users\Abhishek B G\Downloads\Invoice generation samples\RAC_Automation-main\RAC_Automation-main'
OUT_DIR = os.path.dirname(os.path.abspath(__file__))
HISTORY_OUT = os.path.join(OUT_DIR, 'rac_history.json')
DRIVERS_OUT = os.path.join(OUT_DIR, 'rac_drivers.json')

MAP = {
    'SL NO': 'sl_no',
    'BOOKING ID': 'source_booking_id',
    'DATE': 'trip_date',
    'NAME': 'source_name',
    'CAB REG NO': 'source_vehicle_no',
    'MOBIL NO': 'source_mobile',
    'MOBILE NO': 'source_mobile',
    'MOBILE': 'source_mobile',
    'PLAND START': 'planned_start',
    'PLANNED START': 'planned_start',
    'END LOACTION': 'route_text',
    'END LOCATION': 'route_text',
    'PLAND END': 'route_text',
    'PICKUP TIME': 'pickup_time',
    'END TIME': 'end_time',
    'TOTAL HRS SMT': 'total_hours_text',
    'CAB TYPE': 'cab_type',
    'EMP NAME': 'employee_name',
    'START KM': 'start_km',
    'END KM': 'end_km',
    'SMT TOTAL KM': 'total_km',
    'DUTY TYPE': 'duty_type',
    'TOLL': 'toll',
    'PARKING': 'parking',
    'AMOUNT': 'amount',
    'REMARKS': 'remarks',
}


def norm(col):
    return re.sub(r'\s+', ' ', str(col or '').strip().upper())


def clean_text(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    return str(value).replace('\u0000', '').replace('\xa0', ' ').strip()


def num(value, default=0):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return default
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    matches = re.findall(r'\d+(?:\.\d+)?', str(value))
    return float(matches[-1]) if matches else default


def parse_date(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if hasattr(value, 'strftime'):
        year = value.year + 100 if value.year < 2000 else value.year
        return f'{year:04d}-{value.month:02d}-{value.day:02d}'
    text = str(value).strip().replace(' 00:00:00', '')
    if text.lower() in ('', 'nan', 'nat', 'none'):
        return None
    text = text.replace('.', '-')
    for fmt in ('%Y-%m-%d', '%d-%b-%y', '%d-%b-%Y', '%d/%m/%Y', '%d/%m/%y', '%d-%m-%Y', '%d-%m-%y'):
        try:
            parsed = datetime.strptime(text[:20], fmt)
            year = parsed.year + 100 if parsed.year < 2000 else parsed.year
            return f'{year:04d}-{parsed.month:02d}-{parsed.day:02d}'
        except ValueError:
            continue
    return None


def date_from_booking(bid):
    match = re.match(r'^B(\d{2})(\d{2})(\d{2})\b', str(bid or '').strip().upper())
    if not match:
        return None
    year = 2000 + int(match.group(1))
    month = int(match.group(2))
    day = int(match.group(3))
    try:
        datetime(year, month, day)
    except ValueError:
        return None
    return f'{year:04d}-{month:02d}-{day:02d}'


def clean_cab(value):
    text = clean_text(value)
    if not text or re.match(r'^\d', text):
        return ''
    return text.title()


def split_phones(raw):
    digits = re.findall(r'\d{10,12}', re.sub(r'\D+', ' ', clean_text(raw)))
    phones = []
    for item in digits:
        if len(item) == 12 and item.startswith('91'):
            item = item[2:]
        if len(item) == 10:
            phones.append(item)
    unique = []
    for phone in phones:
        if phone not in unique:
            unique.append(phone)
    return unique


def driver_key(name):
    return re.sub(r'\s+', ' ', clean_text(name)).upper()


def find_header(ws):
    for index, row in enumerate(ws.iter_rows(max_row=20, values_only=True), start=1):
        labels = [norm(cell) for cell in row]
        if 'DATE' in labels and 'NAME' in labels:
            fields = [MAP.get(norm(cell), '') for cell in row]
            return index, fields
    return None, []


def parse_sheet(ws, header_row, fields):
    rows = []
    for index, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=1):
        mapped = {}
        empty = True
        for field, val in zip(fields, raw or ()):
            if not field or val in (None, ''):
                continue
            empty = False
            if field not in mapped:
                mapped[field] = val
        if empty:
            continue
        bid = clean_text(mapped.get('source_booking_id'))
        if bid.lower() in ('nan', 'none'):
            bid = ''
        date = date_from_booking(bid) or parse_date(mapped.get('trip_date'))
        name = clean_text(mapped.get('source_name'))
        if not date or not name:
            continue
        if name.upper() in ('NAME', 'SL NO', 'DATE'):
            continue
        rows.append((mapped, bid, date, name, index))
    return rows


def excel_files():
    names = []
    for name in os.listdir(SRC_DIR):
        if name.lower().endswith('.xlsx') and not name.startswith('~$'):
            names.append(os.path.join(SRC_DIR, name))
    return sorted(names)


def build():
    bookings = []
    seen_ids = set()
    seen_duty = set()
    drivers = defaultdict(lambda: {
        'driver_name': '',
        'whatsapp_number': '',
        'alternate_number': '',
        'vehicle_number': '',
        'vehicle_type': '',
        'trip_count': 0,
    })
    stats = Counter()
    sheet_stats = []

    for path in excel_files():
        book = os.path.basename(path)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            print('SKIP FILE', book, exc)
            continue
        for sheet in wb.sheetnames:
            if 'TOLL' in sheet.strip().upper():
                continue
            ws = wb[sheet]
            fields = []
            header_row = 0
            parsed = []
            for index, raw in enumerate(ws.iter_rows(values_only=True), start=1):
                if not fields:
                    labels = [norm(cell) for cell in raw]
                    if 'DATE' in labels and 'NAME' in labels:
                        fields = [MAP.get(norm(cell), '') for cell in raw]
                        header_row = index
                    elif index >= 20:
                        break
                    continue
                mapped = {}
                empty = True
                for field, val in zip(fields, raw or ()):
                    if not field or val in (None, ''):
                        continue
                    empty = False
                    if field not in mapped:
                        mapped[field] = val
                if empty:
                    continue
                bid = clean_text(mapped.get('source_booking_id'))
                if bid.lower() in ('nan', 'none'):
                    bid = ''
                date = date_from_booking(bid) or parse_date(mapped.get('trip_date'))
                name = clean_text(mapped.get('source_name'))
                if not date or not name or name.upper() in ('NAME', 'SL NO', 'DATE'):
                    continue
                parsed.append((mapped, bid, date, name, index - header_row))
            if not fields:
                continue
            kept = 0
            for mapped, bid, date, name, index in parsed:
                if not bid:
                    bid = f"HIST-{date}-{mapped.get('sl_no') or index}-{name[:12]}"
                if bid in seen_ids:
                    bid = f'{bid}-{index}'
                duty_key = (
                    date,
                    driver_key(name),
                    clean_text(mapped.get('source_vehicle_no')),
                    clean_text(mapped.get('pickup_time')),
                )
                if duty_key in seen_duty:
                    continue
                seen_ids.add(bid)
                seen_duty.add(duty_key)
                phones = split_phones(mapped.get('source_mobile'))
                vehicle = clean_text(mapped.get('source_vehicle_no'))
                cab = clean_cab(mapped.get('cab_type'))
                bookings.append({
                    'booking_id': f'hist-{len(bookings):05d}',
                    'source_booking_id': bid,
                    'sl_no': mapped.get('sl_no') if mapped.get('sl_no') == mapped.get('sl_no') else '',
                    'trip_date': date,
                    'source_name': name,
                    'source_vehicle_no': vehicle,
                    'source_mobile': ' / '.join(phones),
                    'planned_start': clean_text(mapped.get('planned_start')),
                    'route_text': clean_text(mapped.get('route_text')),
                    'pickup_time': clean_text(mapped.get('pickup_time')),
                    'end_time': clean_text(mapped.get('end_time')),
                    'total_hours_text': clean_text(mapped.get('total_hours_text')),
                    'cab_type': cab,
                    'employee_name': clean_text(mapped.get('employee_name')) or name,
                    'start_km': num(mapped.get('start_km')),
                    'end_km': num(mapped.get('end_km')),
                    'total_km': num(mapped.get('total_km')),
                    'duty_type': clean_text(mapped.get('duty_type')),
                    'toll': num(mapped.get('toll')),
                    'parking': num(mapped.get('parking')),
                    'amount': num(mapped.get('amount')),
                    'remarks': clean_text(mapped.get('remarks')),
                    'history_source': f'{book} / {sheet}',
                })
                key = driver_key(name)
                row = drivers[key]
                row['driver_name'] = row['driver_name'] or re.sub(r'\s+', ' ', name).strip().title()
                row['trip_count'] += 1
                if phones and not row['whatsapp_number']:
                    row['whatsapp_number'] = '+91' + phones[0]
                if len(phones) > 1 and not row['alternate_number']:
                    row['alternate_number'] = '+91' + phones[1]
                if vehicle:
                    row['vehicle_number'] = vehicle
                if cab:
                    row['vehicle_type'] = cab
                kept += 1
            sheet_stats.append((book, sheet, kept))
            stats['sheets'] += 1
        wb.close()
        stats['files'] += 1

    driver_rows = []
    for index, item in enumerate(sorted(drivers.values(), key=lambda row: row['driver_name'].lower())):
        driver_rows.append({
            'driver_id': f'drv-hist-{index + 1:04d}',
            'driver_name': item['driver_name'],
            'whatsapp_number': item['whatsapp_number'],
            'alternate_number': item['alternate_number'],
            'vehicle_number': item['vehicle_number'],
            'vehicle_type': item['vehicle_type'],
            'home_area': '',
            'active_status': True,
            'total_assignments': item['trip_count'],
            'notes': 'Imported from RAC Excel. Add WhatsApp later.' if not item['whatsapp_number'] else 'Imported from RAC Excel.',
        })

    with open(HISTORY_OUT, 'w', encoding='utf-8') as handle:
        json.dump(bookings, handle, ensure_ascii=False, separators=(',', ':'))
    with open(DRIVERS_OUT, 'w', encoding='utf-8') as handle:
        json.dump(driver_rows, handle, ensure_ascii=False, indent=2)

    print('files', stats['files'], 'duty sheets', stats['sheets'])
    print('bookings', len(bookings), 'kb', round(os.path.getsize(HISTORY_OUT) / 1024, 1))
    print('drivers', len(driver_rows), 'without phone', sum(1 for row in driver_rows if not row['whatsapp_number']))
    print('months', dict(sorted(Counter(row['trip_date'][:7] for row in bookings).items())))
    for book, sheet, kept in sheet_stats:
        print(f'  {kept:5d}  {book} :: {sheet}')


if __name__ == '__main__':
    build()
